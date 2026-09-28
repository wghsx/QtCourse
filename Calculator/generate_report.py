"""Create a truthful, editable lab report from verified local evidence."""

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "实验1-带键盘事件的计算器-实验报告.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal.font.size = Pt(10.5)
zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None:
    zoom.set(qn("w:percent"), "100")

doc.add_heading("实验一：带键盘事件的计算器", 0)
doc.add_paragraph("课程：Qt应用程序开发    学期：2026年秋季")
doc.add_paragraph("姓名：________    学号：________    班级：________")
doc.add_paragraph("实验环境：Windows 11，Qt 6.8.3，MinGW 13.1.0，Qt Creator 20")

doc.add_heading("一、实验目标与界面设计", 1)
doc.add_paragraph("使用 Qt Widgets 制作支持鼠标和键盘输入的四则运算计算器。界面在 Qt Designer 可编辑的 calculatorwindow.ui 中，采用 QVBoxLayout 与 QGridLayout，表达式和结果由两个 QLabel 显示。按钮按数字、运算和控制功能区分配色。")

doc.add_heading("二、主要设计与代码", 1)
doc.add_paragraph("CalculatorWindow 将所有按钮的 clicked 信号连接到 dispatch()；keyPressEvent() 把键盘按键转换成同样的命令，也调用 dispatch()。因此鼠标和键盘共用 CalculatorEngine::input() 的状态处理逻辑。")
doc.add_paragraph("CalculatorEngine 保存第一个操作数 m_left、待执行操作符 m_operator、当前显示值 m_display，以及等待新数字、计算结束和错误状态。输入第二个数后选择新操作符，会先算前一步；连续按操作符则替换待执行操作符。结果显示后按数字开始新计算，按操作符则继续用该结果运算。")
doc.add_paragraph("小数点仅在当前操作数不含 '.' 时加入；退格删除当前输入末位；C 清除全部状态。除数为 0 时显示“不能除以零”，随后可按数字或 C 重新开始。")

doc.add_heading("三、运行结果", 1)
doc.add_paragraph("Windows 图形环境实际运行截图，输入 12.5 + 3.5 =，结果为 16：")
doc.add_picture(str(ROOT / "calculator-screenshot.png"), width=Inches(3.1))

doc.add_heading("四、边界输入与测试", 1)
table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
for cell, title in zip(table.rows[0].cells, ("输入", "处理方法", "验证结果")):
    cell.text = title
for row in (
    ("1..2 - 3 =", "同一操作数只接受一个小数点", "-1.8"),
    ("1 + - 3 =", "连续操作符用后者替换前者", "-2"),
    ("8 / 0 =；再按 2 + 3 =", "除零报错并重置状态", "不能除以零；5"),
    ("12 退格 + 2 =", "退格删除当前操作数末位", "3"),
    ("12.5 + 3.5 =；再按 ×2 -4 /7 =", "连续计算复用前一步结果", "16；4"),
):
    cells = table.add_row().cells
    for cell, value in zip(cells, row): cell.text = value
doc.add_paragraph("自动化测试结果：QtTest 共 7 项通过，0 项失败；包含鼠标点击与键盘混合操作。测试日志见 tests/results-windows.txt。")
doc.add_paragraph("原实验材料仅提供任务模板，未提供可运行的原程序，因此无法实测“原程序现象”。开发中的实际问题是：最初将 role 写成 UI 控件属性，uic 生成不存在的 setRole() 调用导致编译失败；改为按对象名编写样式后通过。另一问题是系统中的其他 g++ 与 Qt 套件不匹配，链接出现 __imp___argc 错误；将 Qt 对应的 MinGW 13.1.0 加到 PATH 前面后构建成功。")

doc.add_heading("五、AI 辅助开发记录", 1)
doc.add_paragraph("问题：怎样让鼠标和键盘操作一致，并处理连续操作及异常输入？提示词：完成带键盘事件的 Qt 计算器，使用 UI 设计器布局、统一处理逻辑并覆盖除零等边界。AI 给出的方案：把按钮和键盘映射为统一命令，在独立的 CalculatorEngine 中维护状态。实际运行：编译、QtTest 和界面截图均已验证。存在的问题与修改：初版 UI 自定义属性及编译工具链配置引起构建失败，已依据编译错误调整 UI 样式和 PATH。")

doc.add_heading("六、Git 记录与实验总结", 1)
doc.add_paragraph("源码和实验记录已推送至 https://github.com/wghsx/QtCourse 的 Calculator/ 目录，保留至少 5 次对应实现、界面、测试、构建说明和报告的提交记录。已录制 2 分 05 秒演示视频，连续展示鼠标小数计算、键盘计算、退格与清除、除零、连续运算、重复小数点与操作符处理，以及实际 Git 提交历史。")
doc.add_paragraph("本实验将输入状态与界面分开，便于检查连续操作、重复小数点等问题。通过 QtTest 核对计算结果，并用 Windows 图形环境截图检查实际界面。")

doc.save(OUT)
print(OUT)
